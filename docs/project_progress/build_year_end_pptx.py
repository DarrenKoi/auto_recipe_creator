# -*- coding: utf-8 -*-
"""Align Tuning Agent 2026년 성과 보고 - PowerPoint(.pptx) 생성 스크립트.

짝 문서 align_tuning_agent_year_end_report_2026.html 과 같은 내용이다. 내용을 고치면 두 파일을
함께 고친다. 한글 폰트는 오피스 Windows 기준 맑은 고딕.

실행:
  uv run python docs/project_progress/build_year_end_pptx.py
출력:
  docs/project_progress/Align_Tuning_Agent_2026_year_end.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt, Emu

HERE = Path(__file__).resolve().parent
OUT = HERE / "Align_Tuning_Agent_2026_year_end.pptx"

FONT = "맑은 고딕"
INK = RGBColor(0x15, 0x11, 0x0D)
MUTED = RGBColor(0x6B, 0x63, 0x5A)
ACCENT = RGBColor(0xC7, 0x5A, 0x3C)
ACCENT_SOFT = RGBColor(0xF3, 0xDC, 0xD2)
FIELD = RGBColor(0x2A, 0x26, 0x22)
FIELD_INK = RGBColor(0xF8, 0xF4, 0xEC)
SURFACE = RGBColor(0xFB, 0xF8, 0xF3)
BORDER = RGBColor(0xE2, 0xDC, 0xD3)
GREY = RGBColor(0xA8, 0xA0, 0x96)


def px(v):
    """1280x720 px 좌표 -> EMU (16:9 = 13.333x7.5in @96dpi)."""
    return Emu(int(v * 9525))


def _font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, lines, size=16, color=INK, bold=False, anchor=MSO_ANCHOR.TOP):
    """lines: str 또는 [str | (str, size, bold, color)]. 줄마다 문단 하나."""
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = px(4)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        s, sz, b, c = (line, size, bold, color) if isinstance(line, str) else line
        run = p.add_run()
        _font(run, sz, b, c)
        run.text = s
    return box


def rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(w), px(h))
    shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def header(slide, no, title):
    text(slide, 60, 36, 1160, 30, f"ALIGN TUNING AGENT · 2026 성과 보고 · {no}", 11, MUTED)
    text(slide, 60, 62, 1160, 50, title, 28, INK, True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(60), px(118), px(1160), px(3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = INK
    bar.line.fill.background()


def metric(slide, x, y, w, h, n, label, hero=False):
    rect(slide, x, y, w, h, FIELD if hero else SURFACE, None if hero else BORDER)
    text(slide, x + 20, y + 18, w - 40, 50, n, 30, FIELD_INK if hero else ACCENT, True)
    text(slide, x + 20, y + 72, w - 40, h - 80, label, 13, FIELD_INK if hero else MUTED)


def card(slide, x, y, w, h, title, bullets, tag=None):
    rect(slide, x, y, w, h, SURFACE, BORDER)
    text(slide, x + 20, y + 14, w - 40, 32, title, 18, INK, True)
    if tag:
        t = rect(slide, x + w - 130, y + 18, 110, 26, ACCENT_SOFT)
        tf = t.text_frame
        tf.margin_top = tf.margin_bottom = 0
        run = tf.paragraphs[0].add_run()
        _font(run, 11, True, ACCENT)
        run.text = tag
    text(slide, x + 20, y + 54, w - 40, h - 64, [f"• {b}" for b in bullets], 13.5)


def table(slide, x, y, w, rows, col_w, size=13, row_h=40):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), px(x), px(y), px(w), px(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = px(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = FIELD if i == 0 else (SURFACE if i % 2 else RGBColor(0xF4, 0xEF, 0xE8))
            cell.margin_left = cell.margin_right = px(10)
            tf = cell.text_frame
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            _font(run, size - 1 if i == 0 else size, i == 0 or j == 0, FIELD_INK if i == 0 else INK)
            run.text = val
    return shp


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1280), px(720)
    blank = prs.slide_layouts[6]

    # 1. 표지
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 1280, 720, FIELD).adjustments[0] = 0
    text(s, 80, 210, 1100, 30, "ALIGN TUNING AGENT · 2026-10-21", 14, GREY)
    text(s, 80, 250, 1100, 80, "Align Tuning Agent 2026년 성과 보고", 40, FIELD_INK, True)
    text(s, 80, 350, 1000, 90,
         "R3 CD-SEM 의 Align Fail 을 사람 대신 감지 · 접속 · 재정렬하는 AI Agent 를 만들어, "
         "2026년에 실장비 완전 자동 제어까지 열었습니다.", 18, GREY)

    # 2. 요약
    s = prs.slides.add_slide(blank)
    header(s, "01", "요약")
    metric(s, 60, 145, 380, 120, "5분 → 1분 미만", "Align Fail 1건당 대응 시간 (실장비 실측)", hero=True)
    metric(s, 450, 145, 380, 120, "31.8 → 76.4%", "Align Key 첫 시도 정답률 (Recipe 200개 오프라인 평가)")
    metric(s, 840, 145, 380, 120, "99%", "화면 버튼 인식 + 클릭 성공률 (500회)")
    table(s, 60, 290, 1160, [
        ["성과", "상태", "내용"],
        ["VLM 인프라 자체 확보", "운영 중", "사내 HCP 에서 5종 평가 -> MAI-UI + PaddleOCR-VL 2종 상시 운영 (외부 API 없음, GPU 1~2장)"],
        ["비표준 화면 GUI 자동화", "실장비 확인", "RCS 로그인 · 장비 접속 · 창 닫기를 coarse -> fine -> OCR 확인 3단계로 자동 수행"],
        ["정렬 정확도 개선", "반영 완료", "성공 이미지 합성 기준(consensus) + OM/SEM 분리 재순위로 첫 시도 정답률 2.4배"],
        ["실시간 무인 대응 루프", "완전 자동 개통", "알람 감지 -> 접속 -> 재정렬 -> OK 확정 -> 기록 -> 종료 (2026-09-15)"],
        ["알람 감지 확산", "활용 중", "같은 알람 감지를 스큐노노 실시간 알람에 적용, 인라인 교대 구성원이 실시간 대응에 활용"],
        ["R3 26대 실전 배치", "준비 중", "AI 전용 RCS 계정 발행 · IP 방화벽 해제 완료, Headless 제어 가능성 테스트 중"],
    ], [250, 150, 760], size=13, row_h=54)

    # 3. 문제
    s = prs.slides.add_slide(blank)
    header(s, "02", "해결한 업무 문제")
    rect(s, 60, 150, 520, 340, SURFACE, BORDER)
    text(s, 84, 166, 480, 24, "이전 · 사람 대응", 12, MUTED)
    text(s, 84, 192, 480, 60, "5분", 40, MUTED, True)
    text(s, 84, 256, 480, 30, "Align Fail 1건당 장비 idle 시간", 15, INK, True)
    text(s, 84, 296, 480, 190, [
        "• 주 400~500건 발생, 교대 구성원이 RCS 모니터를 24시간 감시",
        "• 접속해 배율을 바꿔가며 육안으로 Align Key 탐색",
        "• 실패 화면이 남지 않음",
    ], 14)
    text(s, 590, 290, 100, 60, "→", 40, ACCENT, True)
    rect(s, 700, 150, 520, 340, FIELD)
    text(s, 724, 166, 480, 24, "현재 · AI Agent", 12, GREY)
    text(s, 724, 192, 480, 60, "1분 미만", 40, ACCENT, True)
    text(s, 724, 256, 480, 30, "감지부터 재정렬까지", 15, FIELD_INK, True)
    text(s, 724, 296, 480, 190, [
        "• 10초 주기로 알람(9006) 확인 후 바로 접속",
        "• CV(Computer Vision, 이미지 매칭)가 좌표를 정하고 OK 까지 누름",
        "• 대응 1건마다 녹화 · 이미지 · 기록 1줄",
    ], 14, FIELD_INK)
    text(s, 60, 520, 1160, 120, [
        "계측 화면은 표준 UI 가 아니라 일반 RPA 가 버튼을 못 찾고, Align Key 는 공정 변동으로 매번 달라 "
        "고정 템플릿으로 풀리지 않았습니다.",
        ("VLM = \"어디를 볼지\" · OCR = \"정말 그 버튼인지\" · CV = \"정확히 몇 픽셀인지\" 로 역할을 나눴습니다.",
         15, True, ACCENT),
    ], 15)

    # 4. 올해 만든 것
    s = prs.slides.add_slide(blank)
    header(s, "03", "올해 만든 것")
    cards = [
        ("VLM 인프라", "운영 중", [
            "오픈소스 소형 모델 5종(VLM 3 + OCR 2)을 사내 HCP 에 배포 · 평가",
            "업무 화면 벤치로 이긴 2종만 상시 운영, 모델 교체는 측정으로 결정"]),
        ("GUI 자동화", "실장비 확인", [
            "영역 -> 정밀 좌표 -> OCR 확인으로 바꿔 클릭 성공률 99%",
            "접속 전 Connection User 를 읽어 엔지니어가 쓰는 장비에는 들어가지 않음"]),
        ("정렬 정확도 (CV 벤치)", "반영 완료", [
            "Recipe 200개 오프라인 벤치에서 A/B, 이긴 것만 실시간 루프로 이식",
            "남은 오답의 약 87% 는 등록 위치 문제 -> 재등록 대상 목록 자동 산출"]),
        ("실시간 대응 루프", "완전 자동", [
            "Align Key 가 화면에 없으면 배율을 낮춰 격자로 찾는 Search Around",
            "OK 뒤 다음 측정 위치까지 이어서 보정, Cube 알림은 반드시 전달"]),
    ]
    for i, (t, tag, b) in enumerate(cards):
        card(s, 60 + (i % 2) * 590, 150 + (i // 2) * 270, 570, 250, t, b, tag)

    # 5. 정확도 + 기대 효과
    s = prs.slides.add_slide(blank)
    header(s, "04", "정렬 정확도와 기대 효과")
    cd = CategoryChartData()
    cd.categories = ["첫 시도 정답률", "후보 포함률"]
    cd.add_series("이전 (등록 이미지 1장)", (0.318, 0.434))
    cd.add_series("현재 (consensus)", (0.764, 0.876))
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, px(60), px(140), px(600), px(400), cd).chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    ch.value_axis.maximum_scale = 1.0
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.tick_labels.font.size = Pt(11)
    ch.category_axis.tick_labels.font.size = Pt(13)
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = "0.0%"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(12)
    for ser, col in zip(plot.series, (GREY, ACCENT)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    text(s, 60, 550, 600, 60, "Recipe 200개 오프라인 평가 기준. 현장 자동 복구 성공률은 배치 후 건별 기록으로 집계합니다.", 12, MUTED)
    metric(s, 700, 150, 520, 130, "주 약 27시간", "회수 장비 시간 = Wafer 약 100장/주 추가 측정 (건당 4분 × 주 400건)", hero=True)
    text(s, 700, 300, 520, 300, [
        "• 장비 약 0.2대(Capex 약 4~5억 원)에 해당, 배치 후 실측값으로 갱신",
        "• 야간 · 주말 대응, 실패 화면과 녹화가 남아 원인 분석과 개선 순서를 데이터로 결정",
        "• 버튼 좌표나 이미지를 저장하지 않아, 다른 장비군에도 프롬프트와 기준 이미지만 바꿔 적용",
    ], 14)

    # 6. 이정표
    s = prs.slides.add_slide(blank)
    header(s, "05", "2026 주요 이정표")
    table(s, 60, 150, 1160, [
        ["시기", "내용"],
        ["상반기", "VLM 5종 사내 배포, GUI 자동화 PoC, CV 정확도 벤치 구축과 consensus 채택"],
        ["07월", "OM/SEM 분리 재순위를 운영에 반영 (SEM 첫 시도 정답률 약 +0.07)"],
        ["08월", "모델 2종 상시 운영으로 정리, 실장비 보정 반자동 개통 (재정렬 자동 · OK 는 사람)"],
        ["09월", "List 점유 게이트 오피스 확인 후 완전 자동 제어 전환, 재정렬 closed-loop, "
                 "OM/OM-D/SEM 모드 전환 확인, AX 업무 혁신 공모전 제출"],
        ["10월", "Search Around · 알람 대응 시연 영상 제작, Cube 알림 유실 방지 (디스크 outbox)"],
    ], [160, 1000], size=14, row_h=76)

    # 7. 개발 규모 + 이후 과제
    s = prs.slides.add_slide(blank)
    header(s, "06", "개발 규모와 이후 과제")
    metric(s, 60, 145, 370, 120, "108,294줄", "Python 코드 (실시간 루프 · CV 벤치 · VLM 서빙)", hero=True)
    metric(s, 445, 145, 370, 120, "1,617건", "2026년 커밋")
    metric(s, 830, 145, 390, 120, "1,329개", "실장비 없이 도는 자동 테스트 (09-26 집계)")
    table(s, 60, 290, 560, [
        ["구분", "줄 수"],
        ["실시간 대응 루프", "77,836"],
        ["CV 정확도 벤치", "21,517"],
        ["VLM 서빙 · 배포", "5,855"],
        ["상태 머신 계층", "3,086"],
    ], [340, 220], size=14, row_h=52)
    card(s, 650, 290, 570, 330, "이후 과제", [
        "[테스트 중] reverse engineering 으로 RCS 화면 없이(headless) 장비 제어 -> "
        "GUI 자동화처럼 Agent 가 PC 를 점유하지 않아, 가동 중에도 엔지니어가 다른 작업 병행",
        "AI 전용 계정 · 방화벽 완료, 상시 운영 PC 마무리 -> 단일 장비 pilot -> R3 26대 확대",
        "Search Around · 다음 위치 추적을 실알람에서 검증, 건별 실측 성공률 집계",
        "재등록 대상 Recipe 를 현업과 대조해 Align Key 재등록 추진",
        "2027년: VLM 이 Recipe Tuning 을 직접 수행하도록 확장",
    ])
    text(s, 60, 640, 1160, 40, "최종 목표 · 교대조의 Align Fail 대응을 Agent 가 대신 진행", 18, ACCENT, True)

    prs.save(OUT)
    print(f"[INFO] saved {OUT}")


if __name__ == "__main__":
    build()
